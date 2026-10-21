# 講座「複数プロジェクトの進捗を管理・分析する」の進捗管理表を2つ作る（すべて架空のデータ）
# 使い方: python3 tools/practice/make_progress.py materials/practice
import os, sys, datetime as dt
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

OUT = sys.argv[1]
os.makedirs(OUT, exist_ok=True)
FONT = "Yu Gothic"
NAVY = "1F2B45"
thin = Side(style="thin", color="C9D0DB")
BOX = Border(left=thin, right=thin, top=thin, bottom=thin)
STATUS_FILL = {
    "完了": "D9EAD3",
    "進行中": "DDEBF7",
    "遅延": "F8D7DA",
    "未着手": "F2F2F2",
}
BAR_FILL = {"予定": "BDD7EE", "完了": "93C47D", "進行中": "6FA8DC", "遅延": "E06666"}


def f(bold=False, size=11, color="000000"):
    return Font(name=FONT, bold=bold, size=size, color=color)


def header(ws, row, labels, col=1):
    for i, t in enumerate(labels):
        c = ws.cell(row=row, column=col + i, value=t)
        c.font = f(True, 10, "FFFFFF")
        c.fill = PatternFill("solid", fgColor=NAVY)
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        c.border = BOX


def d(m, day):
    return dt.date(2026, m, day)


def biz_days(a, b):
    # a から b までの営業日（土日を除く）。a の翌日から数える
    n, x = 0, a
    while x < b:
        x += dt.timedelta(days=1)
        if x.weekday() < 5:
            n += 1
    return n


def build(path, pj):
    wb = Workbook()

    # ---------- ダッシュボード ----------
    ws = wb.active
    ws.title = "ダッシュボード"
    for col, w in zip("ABCDEFG", [22, 14, 14, 14, 14, 16, 40]):
        ws.column_dimensions[col].width = w
    ws["A1"] = f"進捗管理表　{pj['name']}"
    ws["A1"].font = f(True, 16)
    ws["A2"] = f"基準日：{pj['today']:%Y/%m/%d}　PM：{pj['pm']}　※ 会社名・人名・数字はすべて架空です"
    ws["A2"].font = f(False, 10, "555555")

    tasks = pj["tasks"]
    count = {s: sum(1 for t in tasks if t[8] == s) for s in STATUS_FILL}
    done_rate = round(sum(t[7] for t in tasks) / len(tasks))

    r = 4
    ws.cell(row=r, column=1, value="全体の状況").font = f(True, 12)
    header(ws, r + 1, ["全タスク数", "完了", "進行中", "遅延", "未着手", "全体進捗率"])
    vals = [len(tasks), count["完了"], count["進行中"], count["遅延"], count["未着手"], done_rate / 100]
    for i, v in enumerate(vals, 1):
        c = ws.cell(row=r + 2, column=i, value=v)
        c.font = f(True, 14)
        c.alignment = Alignment(horizontal="center")
        c.border = BOX
    ws.cell(row=r + 2, column=6).number_format = "0%"
    ws.cell(row=r + 2, column=4).font = f(True, 14, "C00000")

    r = 9
    ws.cell(row=r, column=1, value="区分ごとの件数").font = f(True, 12)
    header(ws, r + 1, ["区分", "完了", "進行中", "遅延", "未着手", "計"])
    cats = []
    for t in tasks:
        if t[1] not in cats:
            cats.append(t[1])
    for i, cat in enumerate(cats):
        row = r + 2 + i
        ws.cell(row=row, column=1, value=cat)
        ts = [t for t in tasks if t[1] == cat]
        for j, s in enumerate(["完了", "進行中", "遅延", "未着手"], 2):
            c = ws.cell(row=row, column=j, value=sum(1 for t in ts if t[8] == s))
            c.alignment = Alignment(horizontal="center")
            if s == "遅延" and c.value:
                c.font = f(True, 11, "C00000")
        ws.cell(row=row, column=6, value=len(ts)).alignment = Alignment(horizontal="center")
        for j in range(1, 7):
            ws.cell(row=row, column=j).border = BOX
            if ws.cell(row=row, column=j).font.name != FONT:
                ws.cell(row=row, column=j).font = f()

    r = r + 3 + len(cats)
    ws.cell(row=r, column=1, value="マイルストーン").font = f(True, 12)
    header(ws, r + 1, ["マイルストーン", "当初の予定日", "今の見込み日", "遅れ（営業日）", "状態"])
    for i, (name, plan, now, state) in enumerate(pj["milestones"]):
        row = r + 2 + i
        late = biz_days(plan, now) if now > plan else 0
        for j, v in enumerate([name, plan, now, late, state], 1):
            c = ws.cell(row=row, column=j, value=v)
            c.font = f(late > 0 and j in (4, 5), 11, "C00000" if late and j in (4, 5) else "000000")
            c.border = BOX
            if isinstance(v, dt.date):
                c.number_format = "yyyy/mm/dd"
            if j > 1:
                c.alignment = Alignment(horizontal="center")

    r = r + 3 + len(pj["milestones"])
    ws.cell(row=r, column=1, value="担当者メモ（定例会議より）").font = f(True, 12)
    for i, m in enumerate(pj["memos"]):
        c = ws.cell(row=r + 1 + i, column=1, value=f"・{m}")
        c.font = f(False, 10)

    # ---------- タスク一覧（ガントチャート） ----------
    g = wb.create_sheet("タスク一覧（ガント）")
    labels = ["No.", "区分", "タスク", "担当", "開始予定", "終了予定", "実績終了", "進捗率", "ステータス", "備考"]
    widths = [5, 14, 30, 10, 11, 11, 11, 8, 10, 36]
    for i, w in enumerate(widths, 1):
        g.column_dimensions[get_column_letter(i)].width = w
    g["A1"] = f"タスク一覧（ガントチャート）　{pj['name']}　基準日 {pj['today']:%Y/%m/%d}"
    g["A1"].font = f(True, 13)
    header(g, 3, labels)
    weeks = []
    w0 = pj["gantt_start"]
    while w0 <= pj["gantt_end"]:
        weeks.append(w0)
        w0 += dt.timedelta(days=7)
    for k, w in enumerate(weeks):
        col = len(labels) + 1 + k
        g.column_dimensions[get_column_letter(col)].width = 5.5
        c = g.cell(row=3, column=col, value=f"{w.month}/{w.day}")
        c.font = f(True, 9, "FFFFFF")
        c.fill = PatternFill("solid", fgColor="E29447" if w <= pj["today"] < w + dt.timedelta(days=7) else NAVY)
        c.alignment = Alignment(horizontal="center", vertical="center")
        c.border = BOX
    for i, t in enumerate(tasks):
        row = 4 + i
        _, cat, name, who, start, end, actual_end, rate, status, note = t
        no = i + 1
        vals = [no, cat, name, who, start, end, actual_end, rate / 100, status, note]
        for j, v in enumerate(vals, 1):
            c = g.cell(row=row, column=j, value=v)
            c.font = f(False, 10)
            c.border = BOX
            if isinstance(v, dt.date):
                c.number_format = "m/d"
                c.alignment = Alignment(horizontal="center")
        g.cell(row=row, column=8).number_format = "0%"
        g.cell(row=row, column=8).alignment = Alignment(horizontal="center")
        sc = g.cell(row=row, column=9)
        sc.fill = PatternFill("solid", fgColor=STATUS_FILL[status])
        sc.alignment = Alignment(horizontal="center")
        if status == "遅延":
            sc.font = f(True, 10, "C00000")
        g.cell(row=row, column=10).alignment = Alignment(wrap_text=True, vertical="top")
        bar = BAR_FILL["予定"] if status == "未着手" else BAR_FILL[status]
        for k, w in enumerate(weeks):
            c = g.cell(row=row, column=len(labels) + 1 + k)
            c.border = BOX
            if start <= w + dt.timedelta(days=6) and w <= end:
                c.fill = PatternFill("solid", fgColor=bar)
    g.freeze_panes = "D4"
    lg = 5 + len(tasks)
    g.cell(row=lg, column=2, value="色の見方：").font = f(True, 9)
    for k, (name, col) in enumerate([("完了", "完了"), ("進行中", "進行中"), ("遅延", "遅延"), ("未着手（予定）", "予定")]):
        c = g.cell(row=lg, column=3 + k * 2, value=name)
        c.font = f(False, 9)
        c.fill = PatternFill("solid", fgColor=BAR_FILL[col])
    wb.save(os.path.join(OUT, path))
    print(path, {s: count[s] for s in count}, "全体", done_rate, "%")


# タスク: (No（使わない。上から順に振る）, 区分, タスク, 担当, 開始予定, 終了予定, 実績終了, 進捗率, ステータス, 備考)
SITE = {
    "name": "コーポレートサイト刷新・展示会PJ",
    "pm": "佐藤",
    "today": d(9, 11),
    "gantt_start": d(7, 6),
    "gantt_end": d(10, 26),
    "milestones": [
        ("要件定義の完了", d(7, 31), d(7, 31), "完了"),
        ("デザイン確定", d(9, 4), d(9, 18), "遅れ"),
        ("新サイト公開", d(10, 9), d(10, 23), "遅れ"),
        ("展示会（ビジネスEXPO秋）", d(10, 21), d(10, 21), "予定どおり"),
    ],
    "memos": [
        "トップページのデザイン案は2回目の差し戻し中。社長の確認待ちで、確定が2週間ずれている",
        "制作会社（ひかりデザイン）から、写真素材の支給待ちで止まっているとの連絡あり",
        "Web・システム担当の伊藤さんが別案件の障害対応にかかりきりで、9月は稼働が半分以下",
        "展示会のブース図面は、主催者の確定版を待っている",
        "9/14 着手予定の「ページの組み込み（コーディング）」は、デザイン確定が前提",
    ],
    "tasks": [
        (1, "企画・要件", "現状サイトのアクセス分析", "佐藤", d(7, 6), d(7, 10), d(7, 10), 100, "完了", ""),
        (0, "企画・要件", "競合他社のサイト調査", "佐藤", d(7, 6), d(7, 10), d(7, 10), 100, "完了", ""),
        (2, "企画・要件", "サイトの目的と目標の整理", "佐藤", d(7, 8), d(7, 17), d(7, 16), 100, "完了", ""),
        (3, "企画・要件", "サイトマップ（ページ構成）作成", "鈴木", d(7, 13), d(7, 24), d(7, 24), 100, "完了", ""),
        (4, "企画・要件", "要件定義書の作成と承認", "佐藤", d(7, 20), d(7, 31), d(7, 31), 100, "完了", ""),
        (5, "デザイン制作", "ワイヤーフレーム（画面の骨組み）", "鈴木", d(8, 3), d(8, 12), d(8, 12), 100, "完了", ""),
        (6, "デザイン制作", "トップページのデザイン", "ひかりデザイン", d(8, 10), d(8, 21), None, 70, "遅延", "社長確認で2回目の差し戻し。確定は早くて9/18"),
        (7, "デザイン制作", "下層ページのデザイン", "ひかりデザイン", d(8, 17), d(8, 28), None, 40, "遅延", "トップページ確定待ち"),
        (8, "デザイン制作", "写真撮影・素材の手配", "鈴木", d(8, 17), d(8, 28), None, 50, "遅延", "社内の写真素材の支給が遅れ、制作会社が待っている"),
        (9, "デザイン制作", "ロゴ・色のガイド整理", "鈴木", d(8, 3), d(8, 7), d(8, 7), 100, "完了", ""),
        (10, "Web・システム", "サーバーとドメインの準備", "伊藤", d(8, 3), d(8, 14), d(8, 14), 100, "完了", ""),
        (11, "Web・システム", "CMS（更新の仕組み）の設定", "伊藤", d(8, 24), d(9, 4), None, 60, "遅延", "担当が別案件の障害対応で稼働が半分以下"),
        (12, "Web・システム", "問い合わせフォームの作成", "伊藤", d(8, 31), d(9, 9), None, 20, "遅延", "同上。CMS の設定が終わってから"),
        (13, "Web・システム", "ページの組み込み（コーディング）", "伊藤", d(9, 14), d(10, 2), None, 0, "未着手", "デザイン確定が前提。9/14 着手は難しい見込み"),
        (14, "Web・システム", "表示の確認とテスト", "伊藤", d(10, 5), d(10, 8), None, 0, "未着手", ""),
        (15, "原稿", "会社案内ページの原稿", "高橋", d(8, 3), d(8, 21), d(8, 20), 100, "完了", ""),
        (16, "原稿", "サービス紹介ページの原稿", "高橋", d(8, 17), d(9, 4), d(9, 3), 100, "完了", ""),
        (17, "原稿", "採用ページの原稿", "高橋", d(8, 31), d(9, 18), None, 50, "進行中", ""),
        (18, "展示会準備", "出展の申し込み", "渡辺", d(7, 6), d(7, 10), d(7, 9), 100, "完了", ""),
        (19, "展示会準備", "ブースの設計", "渡辺", d(8, 17), d(9, 4), None, 60, "遅延", "主催者の確定図面を待っている"),
        (20, "展示会準備", "パンフレットの制作", "渡辺", d(8, 24), d(9, 11), None, 30, "遅延", "新サイトのデザインに合わせるため、デザイン確定待ち"),
        (21, "展示会準備", "ノベルティ（配る記念品）の手配", "山本", d(8, 24), d(9, 4), d(9, 2), 100, "完了", ""),
        (22, "展示会準備", "来場のお客様への案内メール", "山本", d(9, 7), d(9, 25), None, 40, "進行中", ""),
        (23, "展示会準備", "当日のスタッフ配置", "山本", d(9, 7), d(9, 18), None, 30, "進行中", ""),
        (24, "展示会準備", "展示会の直前リハーサル", "渡辺", d(10, 13), d(10, 16), None, 0, "未着手", ""),
    ],
}

PROMO = {
    "name": "新商品プロモーションPJ",
    "pm": "中村",
    "today": d(9, 11),
    "gantt_start": d(7, 27),
    "gantt_end": d(11, 9),
    "milestones": [
        ("販促プランの決定", d(8, 7), d(8, 7), "完了"),
        ("広告の素材の完成", d(9, 18), d(9, 25), "遅れ"),
        ("予約の受付開始", d(10, 1), d(10, 1), "予定どおり"),
        ("発売日", d(11, 2), d(11, 2), "予定どおり"),
    ],
    "memos": [
        "インフルエンサー2名のうち1名の契約書が先方の法務チェック中で、返事待ち",
        "広告の動画は、商品の最終サンプルが届かないと撮影できない（工場の出荷が9/15に遅れる）",
        "予約ページは予定どおり。10/1 の受付開始には間に合う見込み",
    ],
    "tasks": [
        (1, "企画", "ターゲットと売り文句の整理", "中村", d(7, 27), d(8, 5), d(8, 4), 100, "完了", ""),
        (2, "企画", "販促プランと予算の決定", "中村", d(8, 3), d(8, 7), d(8, 7), 100, "完了", ""),
        (3, "企画", "発売日までのスケジュール作成", "中村", d(8, 5), d(8, 7), d(8, 7), 100, "完了", ""),
        (4, "広告制作", "広告のキャッチコピー", "小林", d(8, 10), d(8, 21), d(8, 21), 100, "完了", ""),
        (5, "広告制作", "広告の画像（バナー）制作", "小林", d(8, 17), d(9, 4), d(9, 4), 100, "完了", ""),
        (6, "広告制作", "広告の動画の撮影", "小林", d(9, 7), d(9, 11), None, 0, "遅延", "商品の最終サンプル待ち（工場の出荷が9/15に遅れる）"),
        (7, "広告制作", "広告の動画の編集", "小林", d(9, 14), d(9, 18), None, 0, "未着手", "撮影が終わってから"),
        (8, "SNS・インフルエンサー", "インフルエンサーの候補選び", "加藤", d(8, 10), d(8, 21), d(8, 20), 100, "完了", ""),
        (9, "SNS・インフルエンサー", "インフルエンサーとの契約", "加藤", d(8, 24), d(9, 4), None, 50, "遅延", "2名のうち1名が先方の法務チェック中"),
        (10, "SNS・インフルエンサー", "SNS投稿のカレンダー作成", "加藤", d(8, 31), d(9, 11), None, 70, "進行中", ""),
        (11, "SNS・インフルエンサー", "発売前のティザー投稿", "加藤", d(9, 21), d(10, 30), None, 0, "未着手", ""),
        (12, "販売準備", "予約ページの作成", "吉田", d(8, 24), d(9, 18), None, 60, "進行中", ""),
        (13, "販売準備", "店頭ポップの制作", "吉田", d(8, 31), d(9, 11), None, 40, "遅延", "動画の画面写真を使う予定のため、撮影待ち"),
        (14, "販売準備", "販売店への案内資料", "吉田", d(8, 17), d(8, 28), d(8, 27), 100, "完了", ""),
        (15, "販売準備", "問い合わせ対応のQ&A作成", "吉田", d(9, 7), d(9, 25), None, 30, "進行中", ""),
        (16, "販売準備", "予約の受付開始", "吉田", d(10, 1), d(10, 1), None, 0, "未着手", ""),
        (17, "効果測定", "効果を見る指標の決定", "中村", d(8, 10), d(8, 14), d(8, 14), 100, "完了", ""),
        (18, "効果測定", "発売後の振り返り", "中村", d(11, 2), d(11, 6), None, 0, "未着手", ""),
    ],
}

build("progress-corporate-site.xlsx", SITE)
build("progress-new-product.xlsx", PROMO)
