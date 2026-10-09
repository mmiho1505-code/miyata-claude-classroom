# 練習用データ一式を作る（すべて架空のデータ）
import csv, os, random, sys, datetime as dt
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.chart import BarChart, Reference
from openpyxl.utils import get_column_letter

OUT = sys.argv[1]
os.makedirs(OUT, exist_ok=True)
FONT = "Yu Gothic"
NAVY = "1F2B45"
thin = Side(style="thin", color="C9D0DB")
BOX = Border(left=thin, right=thin, top=thin, bottom=thin)
INPUT_FILL = PatternFill("solid", fgColor="FFF7D6")


def base(ws, widths):
    for i, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = w


def f(bold=False, size=11, color="000000"):
    return Font(name=FONT, bold=bold, size=size, color=color)


def header(ws, row, labels, col=1):
    for i, t in enumerate(labels):
        c = ws.cell(row=row, column=col + i, value=t)
        c.font = f(True, 11, "FFFFFF")
        c.fill = PatternFill("solid", fgColor=NAVY)
        c.alignment = Alignment(horizontal="center", vertical="center")
        c.border = BOX


def style_range(ws, r1, r2, c1, c2):
    for r in range(r1, r2 + 1):
        for c in range(c1, c2 + 1):
            cell = ws.cell(row=r, column=c)
            if cell.font is None or cell.font.name != FONT:
                cell.font = f()
            cell.border = BOX


def write_csv(name, head, rows):
    with open(os.path.join(OUT, name), "w", newline="", encoding="utf-8-sig") as fp:
        w = csv.writer(fp)
        w.writerow(head)
        w.writerows(rows)


# ---------- 1. 取引先リスト clients.xlsx ----------
CLIENTS = [
    (1, "株式会社 見本工務店", "総務部 山田 様", "100-0001", "東京都千代田区見本1-2-3"),
    (2, "有限会社 サンプル食品", "経理 佐藤 様", "100-0002", "東京都千代田区例題4-5-6"),
    (3, "合同会社 例題デザイン", "代表 鈴木 様", "100-0003", "東京都千代田区練習7-8-9"),
    (4, "株式会社 架空運送", "管理部 高橋 様", "100-0004", "東京都千代田区架空10-11"),
    (5, "株式会社 ためし商店", "店長 田中 様", "100-0005", "東京都千代田区試作12-13"),
]
LINES = [
    (1, "月次顧問料", 1, 50000, 0.10),
    (1, "資料作成費", 2, 15000, 0.10),
    (2, "月次顧問料", 1, 30000, 0.10),
    (2, "試食会用 弁当（軽減税率）", 12, 800, 0.08),
    (3, "月次顧問料", 1, 30000, 0.10),
    (4, "月次顧問料", 1, 50000, 0.10),
    (4, "研修（半日）", 1, 40000, 0.10),
    (5, "スポット相談", 3, 10000, 0.10),
]
wb = Workbook()
ws = wb.active
ws.title = "取引先"
ws["A1"] = "取引先リスト（練習用・すべて架空の会社です）"
ws["A1"].font = f(True, 14)
header(ws, 3, ["取引先No", "会社名（正式名）", "部署・担当者", "郵便番号", "住所", "締め日", "支払期限"])
for i, (no, name, person, zipc, addr) in enumerate(CLIENTS):
    r = 4 + i
    for c, v in enumerate([no, name, person, zipc, addr, "月末", "翌月末"], 1):
        ws.cell(row=r, column=c, value=v)
style_range(ws, 4, 3 + len(CLIENTS), 1, 7)
base(ws, [11, 26, 20, 12, 32, 9, 11])

ws2 = wb.create_sheet("9月分の明細")
ws2["A1"] = "2026年9月分の請求明細（練習用・架空）"
ws2["A1"].font = f(True, 14)
ws2["A2"] = "金額（税抜）は 数量×単価 の式です。税率は 10% か 8%（軽減税率）。"
ws2["A2"].font = f(False, 10, "5B6577")
header(ws2, 4, ["取引先No", "会社名", "品目", "数量", "単価", "税率", "金額（税抜）"])
for i, (no, item, qty, price, rate) in enumerate(LINES):
    r = 5 + i
    ws2.cell(row=r, column=1, value=no)
    ws2.cell(row=r, column=2, value=f"=INDEX(取引先!$B$4:$B$8,MATCH(A{r},取引先!$A$4:$A$8,0))")
    ws2.cell(row=r, column=3, value=item)
    ws2.cell(row=r, column=4, value=qty)
    ws2.cell(row=r, column=5, value=price).number_format = "#,##0"
    ws2.cell(row=r, column=6, value=rate).number_format = "0%"
    ws2.cell(row=r, column=7, value=f"=D{r}*E{r}").number_format = "#,##0"
last = 4 + len(LINES)
ws2.cell(row=last + 1, column=6, value="合計")
ws2.cell(row=last + 1, column=7, value=f"=SUM(G5:G{last})").number_format = "#,##0"
style_range(ws2, 5, last + 1, 1, 7)
ws2.cell(row=last + 1, column=6).font = f(True)
ws2.cell(row=last + 1, column=7).font = f(True)
base(ws2, [11, 26, 28, 8, 11, 8, 14])
wb.save(os.path.join(OUT, "clients.xlsx"))

# ---------- 2. 請求書ひな形 invoice_template.xlsx ----------
wb = Workbook()
ws = wb.active
ws.title = "請求書"
base(ws, [32, 8, 12, 14, 18])
ws.merge_cells("A1:E1")
ws["A1"] = "請 求 書"
ws["A1"].font = f(True, 20)
ws["A1"].alignment = Alignment(horizontal="center")
ws.row_dimensions[1].height = 34
ws["A3"] = "（宛名：正式な会社名）"
ws["A3"].font = f(True, 13)
ws["A3"].fill = INPUT_FILL
ws["B3"] = "御中"
ws["D3"] = "請求書番号"
ws["E3"] = "2026-09-000"
ws["E3"].fill = INPUT_FILL
ws["D4"] = "発行日"
ws["E4"] = dt.date(2026, 9, 30)
ws["E4"].number_format = "yyyy/m/d"
ws["E4"].fill = INPUT_FILL
ws["D5"] = "支払期限"
ws["E5"] = dt.date(2026, 10, 31)
ws["E5"].number_format = "yyyy/m/d"
ws["E5"].fill = INPUT_FILL
ws["A5"] = "下記のとおりご請求申し上げます。"
ws["A7"] = "ご請求金額（税込）"
ws["A7"].font = f(True, 12)
ws["B7"] = "=E27"
ws["B7"].number_format = '"¥"#,##0'
ws["B7"].font = f(True, 14)
ws.merge_cells("B7:C7")
for r_, t_ in ((8, "株式会社 練習商事（架空）"), (9, "東京都千代田区練習1-1-1"), (10, "TEL 03-0000-0000"), (11, "登録番号 T1234567890123")):
    ws.merge_cells(f"C{r_}:E{r_}")
    ws[f"C{r_}"] = t_
    ws[f"C{r_}"].alignment = Alignment(horizontal="right")
ws["C8"].font = f(True)
header(ws, 13, ["品目", "数量", "単価", "税率", "金額（税抜）"])
EX = [("（例）月次顧問料", 1, 50000, 0.10)]
for i in range(10):
    r = 14 + i
    if i < len(EX):
        item, qty, price, rate = EX[i]
        ws.cell(row=r, column=1, value=item)
        ws.cell(row=r, column=2, value=qty)
        ws.cell(row=r, column=3, value=price)
        ws.cell(row=r, column=4, value=rate)
    for c in (1, 2, 3, 4):
        ws.cell(row=r, column=c).fill = INPUT_FILL
    ws.cell(row=r, column=3).number_format = "#,##0"
    ws.cell(row=r, column=4).number_format = "0%"
    ws.cell(row=r, column=5, value=f'=IF(OR(B{r}="",C{r}=""),"",B{r}*C{r})').number_format = "#,##0"
style_range(ws, 14, 23, 1, 5)
rows = [
    (24, "10%対象（税抜）", "=SUMIF(D14:D23,0.1,E14:E23)"),
    (25, "消費税（10%）", "=ROUNDDOWN(E24*0.1,0)"),
    (26, "8%対象（税抜）／消費税（8%）", None),
    (27, "合計（税込）", "=E24+E25+C26+E26"),
]
ws["A24"], ws["E24"] = rows[0][1], rows[0][2]
ws["A25"], ws["E25"] = rows[1][1], rows[1][2]
ws["A26"] = "8%対象（税抜）"
ws["C26"] = "=SUMIF(D14:D23,0.08,E14:E23)"
ws["D26"] = "消費税（8%）"
ws["E26"] = "=ROUNDDOWN(C26*0.08,0)"
ws["A27"], ws["E27"] = rows[3][1], rows[3][2]
for r in (24, 25, 26, 27):
    for c in (3, 5):
        ws.cell(row=r, column=c).number_format = "#,##0"
style_range(ws, 24, 27, 1, 5)
ws["A27"].font = f(True, 12)
ws["E27"].font = f(True, 12)
ws["A29"] = "お振込先"
ws["A29"].font = f(True)
ws["A30"] = "練習銀行 本店 普通 1234567 カ）レンシュウショウジ"
for r_, t_ in ((32, "【使い方】黄色のセルだけ書き換えます。金額・消費税・合計は式で自動計算です。"), (33, "14行目の（例）は消して使います。"), (34, "会社名・登録番号・振込先は、練習用の架空のものです。"), (35, "自社のものに直してから使ってください。")):
    ws[f"A{r_}"] = t_
    ws[f"A{r_}"].font = f(False, 10, "5B6577")
for row in ws.iter_rows(min_row=3, max_row=11):
    for c in row:
        if c.font is None or c.font.name != FONT:
            c.font = f()
ws.print_area = "A1:E35"
ws.page_setup.orientation = "portrait"
ws.page_setup.paperSize = 9
ws.page_setup.fitToWidth = 1
ws.page_setup.fitToHeight = 1
ws.sheet_properties.pageSetUpPr.fitToPage = True
wb.save(os.path.join(OUT, "invoice_template.xlsx"))

# ---------- 3. 売上データ（架空のアパレル店） ----------
STORES = ["駅前店", "本町店", "モール店"]
STAFF = {"駅前店": ["佐藤", "鈴木"], "本町店": ["高橋", "田中"], "モール店": ["伊藤", "渡辺"]}
PRODUCTS = [
    ("シャツ（白）", "トップス", 3900), ("カットソー", "トップス", 2900), ("ニット", "トップス", 5900),
    ("パーカー", "トップス", 4900), ("デニムパンツ", "ボトムス", 6900), ("チノパンツ", "ボトムス", 5900),
    ("スカート", "ボトムス", 4900), ("ジャケット", "アウター", 12900), ("コート", "アウター", 19800),
    ("ストール", "小物", 2900), ("ベルト", "小物", 3500), ("トートバッグ", "小物", 4500),
]
# 月・店舗・商品ごとの売れ方（重み）。当月はモール店が伸び、本町店はデニムが落ちる
BASE_W = {"駅前店": 1.0, "本町店": 0.8, "モール店": 0.9}
MONTHS = {
    "2025-09": {"days": 30, "n": 430, "store": {"駅前店": 1.0, "本町店": 1.0, "モール店": 0.85}, "prod": {}},
    "2026-08": {"days": 31, "n": 400, "store": {"駅前店": 0.95, "本町店": 0.95, "モール店": 1.0}, "prod": {"コート": 0.2, "ニット": 0.4, "ジャケット": 0.5}},
    "2026-09": {"days": 30, "n": 450, "store": {"駅前店": 1.0, "本町店": 0.78, "モール店": 1.25}, "prod": {"パーカー": 1.6, "ジャケット": 1.3}},
}
DENIM_DROP = {("2026-09", "本町店", "デニムパンツ"): 0.3}
PROD_W = {"シャツ（白）": 1.2, "カットソー": 1.4, "ニット": 0.9, "パーカー": 1.0, "デニムパンツ": 1.1, "チノパンツ": 0.8,
          "スカート": 0.9, "ジャケット": 0.5, "コート": 0.25, "ストール": 0.6, "ベルト": 0.5, "トートバッグ": 0.6}
SALES = {}
for ym, cfg in MONTHS.items():
    rnd = random.Random("sales-" + ym)
    y, m = map(int, ym.split("-"))
    combos, weights = [], []
    for s in STORES:
        for (p, cat, price) in PRODUCTS:
            w = BASE_W[s] * cfg["store"][s] * PROD_W[p] * cfg["prod"].get(p, 1.0) * DENIM_DROP.get((ym, s, p), 1.0)
            combos.append((s, p, cat, price))
            weights.append(w)
    rows = []
    for _ in range(cfg["n"]):
        s, p, cat, price = rnd.choices(combos, weights)[0]
        day = rnd.randint(1, cfg["days"])
        qty = rnd.choices([1, 2, 3], [0.78, 0.18, 0.04])[0]
        rows.append([f"{y}-{m:02d}-{day:02d}", s, p, cat, rnd.choice(STAFF[s]), qty, price, qty * price])
    rows.sort(key=lambda r: (r[0], r[1], r[2]))
    SALES[ym] = rows
    write_csv(f"sales-{ym}.csv", ["日付", "店舗", "商品", "カテゴリ", "販売員", "数量", "単価", "金額"], rows)


def store_total(ym):
    t = {s: 0 for s in STORES}
    for r in SALES[ym]:
        t[r[1]] += r[7]
    return t


prev_year = store_total("2025-09")
targets = {s: int(round(prev_year[s] * 1.05, -4)) for s in STORES}
write_csv("sales-target-2026-09.csv", ["店舗", "目標金額"], [[s, targets[s]] for s in STORES])

# 売上レポートのひな形（グラフ付き）
wb = Workbook()
ws = wb.active
ws.title = "レポート"
base(ws, [16, 14, 14, 12, 14, 12])
ws["A1"] = "月次売上レポート"
ws["A1"].font = f(True, 16)
ws["A2"] = "対象月"
ws["B2"] = "（例）2026年9月"
ws["B2"].fill = INPUT_FILL
header(ws, 4, ["店舗", "売上", "目標", "達成率", "前月売上", "前月比"])
for i, s in enumerate(STORES):
    r = 5 + i
    ws.cell(row=r, column=1, value=s)
    ws.cell(row=r, column=2, value=f"=SUMIFS(明細!$H:$H,明細!$B:$B,A{r})").number_format = "#,##0"
    ws.cell(row=r, column=3).number_format = "#,##0"
    ws.cell(row=r, column=3).fill = INPUT_FILL
    ws.cell(row=r, column=4, value=f'=IF(N(C{r})=0,"",B{r}/C{r})').number_format = "0.0%"
    ws.cell(row=r, column=5).number_format = "#,##0"
    ws.cell(row=r, column=5).fill = INPUT_FILL
    ws.cell(row=r, column=6, value=f'=IF(N(E{r})=0,"",B{r}/E{r})').number_format = "0.0%"
ws["A8"] = "合計"
ws["B8"] = "=SUM(B5:B7)"
ws["C8"] = "=SUM(C5:C7)"
ws["D8"] = '=IF(C8=0,"",B8/C8)'
ws["E8"] = "=SUM(E5:E7)"
ws["F8"] = '=IF(E8=0,"",B8/E8)'
for c, fmt in ((2, "#,##0"), (3, "#,##0"), (4, "0.0%"), (5, "#,##0"), (6, "0.0%")):
    ws.cell(row=8, column=c).number_format = fmt
style_range(ws, 5, 8, 1, 6)
for c in range(1, 7):
    ws.cell(row=8, column=c).font = f(True)
header(ws, 11, ["カテゴリ", "売上", "構成比"])
CATS = ["トップス", "ボトムス", "アウター", "小物"]
for i, cat in enumerate(CATS):
    r = 12 + i
    ws.cell(row=r, column=1, value=cat)
    ws.cell(row=r, column=2, value=f"=SUMIFS(明細!$H:$H,明細!$D:$D,A{r})").number_format = "#,##0"
    ws.cell(row=r, column=3, value=f'=IF($B$8=0,"",B{r}/$B$8)').number_format = "0.0%"
style_range(ws, 12, 15, 1, 3)
ws["A17"] = "ひとことコメント"
ws["A17"].font = f(True)
ws["A18"] = "（ここに今月の所見を書く）"
ws["A18"].fill = INPUT_FILL
ws.merge_cells("A18:F20")
ws["A18"].alignment = Alignment(vertical="top", wrap_text=True)
chart = BarChart()
chart.type = "col"
chart.title = "店舗別 売上と目標"
chart.y_axis.title = "円"
chart.add_data(Reference(ws, min_col=2, min_row=4, max_col=3, max_row=7), titles_from_data=True)
chart.set_categories(Reference(ws, min_col=1, min_row=5, max_row=7))
chart.height, chart.width = 8.5, 15
ws.add_chart(chart, "A22")

wd = wb.create_sheet("明細")
header(wd, 1, ["日付", "店舗", "商品", "カテゴリ", "販売員", "数量", "単価", "金額"])
base(wd, [12, 10, 16, 10, 9, 7, 9, 11])
wu = wb.create_sheet("使い方")
base(wu, [96])
for i, t in enumerate([
    "【このひな形の使い方】（練習用）",
    "1. 「明細」シートの2行目から、その月の売上明細（日付・店舗・商品・カテゴリ・販売員・数量・単価・金額）を入れます。",
    "2. 「レポート」シートの売上・達成率・前月比・カテゴリ別は、式で自動計算されます。式は消さないでください。",
    "3. 黄色のセル（対象月・目標・前月売上・コメント）だけ、手で入れます。",
    "4. グラフは、店舗別の売上と目標の表を参照しています。",
    "店舗名・カテゴリ名は、「明細」と「レポート」で同じ文字にそろえます（駅前店・本町店・モール店／トップス・ボトムス・アウター・小物）。",
    "データはすべて架空です。",
]):
    wu.cell(row=1 + i, column=1, value=t).font = f(i == 0, 12 if i == 0 else 11)
wb.save(os.path.join(OUT, "sales-report-template.xlsx"))

# ---------- 4. アンケート survey.csv ----------
rnd = random.Random("survey")
AGES = ["20代", "30代", "40代", "50代", "60代以上"]
FREQ = ["はじめて", "2〜3回目", "月に1回くらい", "週に1回以上"]
SAT = ["とても満足", "満足", "ふつう", "やや不満", "不満"]
AGAIN = ["ぜひ利用したい", "機会があれば", "わからない", "利用しない"]
HOW = ["知人の紹介", "店の前を通って", "SNS", "チラシ", "検索"]
GOOD = ["スタッフの説明が分かりやすかった", "店内がきれいで落ち着いた", "待ち時間が短かった", "値段がはっきりしていて安心した",
        "相談しやすい雰囲気だった", "予約がかんたんだった", "仕上がりに満足している", "子ども連れでも入りやすかった"]
BAD = ["待ち時間が長かった", "駐車場が分かりにくい", "料金の説明がもう少しほしい", "予約が取りにくい日があった",
       "電話がつながりにくかった", "店内が少し寒かった", "支払い方法を増やしてほしい", "営業時間をもう少し長くしてほしい"]
rows = []
for i in range(1, 61):
    sat = rnd.choices(SAT, [0.3, 0.38, 0.18, 0.1, 0.04])[0]
    si = SAT.index(sat)
    again = rnd.choices(AGAIN, [[0.7, 0.25, 0.05, 0], [0.4, 0.5, 0.1, 0], [0.1, 0.5, 0.35, 0.05], [0.02, 0.3, 0.4, 0.28], [0, 0.1, 0.3, 0.6]][si])[0]
    if si <= 1:
        free = rnd.choice(GOOD) if rnd.random() < 0.75 else ""
        if free and rnd.random() < 0.25:
            free += "。ただ、" + rnd.choice(BAD)
    elif si == 2:
        free = rnd.choice(BAD) if rnd.random() < 0.6 else ""
    else:
        free = rnd.choice(BAD)
    day = rnd.randint(1, 30)
    rows.append([i, f"2026-09-{day:02d}", rnd.choices(AGES, [0.12, 0.22, 0.26, 0.24, 0.16])[0], rnd.choices(FREQ, [0.3, 0.3, 0.28, 0.12])[0],
                 sat, again, rnd.choices(HOW, [0.3, 0.2, 0.22, 0.1, 0.18])[0], free])
rows.sort(key=lambda r: r[1])
for i, r in enumerate(rows, 1):
    r[0] = i
write_csv("survey.csv", ["回答No", "回答日", "年代", "ご利用回数", "満足度", "また利用したいですか", "当店を知ったきっかけ", "ご意見・ご感想（自由記述）"], rows)

# ---------- 5. ABC分析用 sales.xlsx ----------
ABC = [("A-001", "定番シャツ", 412, 3900), ("A-002", "カットソー", 388, 2900), ("A-003", "デニムパンツ", 201, 6900),
       ("A-004", "ニット", 164, 5900), ("A-005", "ジャケット", 58, 12900), ("A-006", "チノパンツ", 96, 5900),
       ("A-007", "パーカー", 102, 4900), ("A-008", "コート", 21, 19800), ("A-009", "スカート", 77, 4900),
       ("A-010", "トートバッグ", 61, 4500), ("A-011", "ストール", 66, 2900), ("A-012", "ベルト", 43, 3500),
       ("A-013", "靴下3足組", 118, 1000), ("A-014", "ハンカチ", 96, 800), ("A-015", "帽子", 22, 2900),
       ("A-016", "手袋", 14, 2500), ("A-017", "ネクタイ", 9, 3900), ("A-018", "エコバッグ", 31, 900),
       ("A-019", "キーケース", 6, 3200), ("A-020", "ピンバッジ", 12, 600)]
wb = Workbook()
ws = wb.active
ws.title = "売上"
ws["A1"] = "商品別売上（2026年4月〜9月・練習用の架空データ）"
ws["A1"].font = f(True, 13)
header(ws, 3, ["商品コード", "商品名", "数量", "単価", "売上金額"])
for i, (code, name, qty, price) in enumerate(ABC):
    r = 4 + i
    ws.cell(row=r, column=1, value=code)
    ws.cell(row=r, column=2, value=name)
    ws.cell(row=r, column=3, value=qty).number_format = "#,##0"
    ws.cell(row=r, column=4, value=price).number_format = "#,##0"
    ws.cell(row=r, column=5, value=f"=C{r}*D{r}").number_format = "#,##0"
last = 3 + len(ABC)
ws.cell(row=last + 1, column=4, value="合計").font = f(True)
ws.cell(row=last + 1, column=5, value=f"=SUM(E4:E{last})").number_format = "#,##0"
style_range(ws, 4, last + 1, 1, 5)
ws.cell(row=last + 1, column=4).font = f(True)
ws.cell(row=last + 1, column=5).font = f(True)
base(ws, [12, 18, 9, 10, 14])
wb.save(os.path.join(OUT, "sales.xlsx"))

# ---------- 6. SNS posts.csv / sample_posts.txt ----------
rnd = random.Random("sns")
THEMES = [
    ("資金繰り", ["月末に慌てないコツは、入金と支払いの日を1枚の表に並べることです。", "「黒字なのにお金がない」は、売掛金の回収が遅いサインかもしれません。", "資金繰り表は、3か月先まで書くだけで見える景色が変わります。"]),
    ("銀行", ["銀行には、決算のときだけでなく毎月顔を出す。それだけで話が早くなります。", "試算表は、頼まれる前に持っていく。小さな習慣が信用になります。", "借りる話より先に、いまの数字を見せる。順番を変えると反応が変わります。"]),
    ("数字の見方", ["粗利は「売上−原価」。まず自社の粗利率を言えるようにしましょう。", "売上が増えたのに利益が減った。そんなときは固定費を1行ずつ見ます。", "前年同月と比べる。前月と比べるより、季節のクセが消えます。"]),
    ("AI活用", ["AIに頼むときは、目的・前提・形式の3つを書く。これだけで答えが変わります。", "議事録はAIの下書き、最後は人の目。決定事項と担当者だけは自分で確かめます。", "AIに数字を作らせない。無い数字は「不明」と書かせるのがコツです。"]),
    ("日々のこと", ["今朝は現場まわり。数字の向こうに、働く人の顔があります。", "今日は月初の締め。机の上の紙が少し減りました。", "お客様からうれしい報告をいただきました。続けてきてよかったです。"]),
]
TAGS = {"資金繰り": "#資金繰り #中小企業", "銀行": "#銀行対応 #経営者", "数字の見方": "#決算書 #数字の見方", "AI活用": "#AI活用 #業務改善", "日々のこと": ""}
posts = []
used = set()
for i in range(40):
    theme, texts = rnd.choice(THEMES)
    text = rnd.choice(texts)
    q = rnd.random() < 0.3
    body = text + ("　あなたの会社ではどうしていますか？" if q else "")
    tag = TAGS[theme] if rnd.random() < 0.7 else ""
    full = (body + (" " + tag if tag else "")).strip()
    media = rnd.choice(["X", "X", "Instagram"])
    day = rnd.randint(1, 30)
    hour = rnd.choice([7, 7, 8, 12, 12, 18, 20, 21, 21, 22])
    base_like = {"資金繰り": 46, "銀行": 38, "数字の見方": 30, "AI活用": 52, "日々のこと": 14}[theme]
    mult = (1.35 if hour in (7, 8, 21) else 1.0) * (1.25 if q else 1.0) * (1.15 if tag else 1.0) * (1.3 if media == "Instagram" else 1.0)
    likes = max(1, int(rnd.gauss(base_like * mult, base_like * 0.25)))
    saves = max(0, int(likes * rnd.uniform(0.08, 0.3) * (1.5 if theme in ("資金繰り", "AI活用") else 0.6)))
    comments = max(0, int(likes * rnd.uniform(0.0, 0.08) * (2.2 if q else 1.0)))
    posts.append([f"2026-09-{day:02d} {hour:02d}:{rnd.choice(['00','05','10','30','45'])}", media, full, likes, saves, comments])
posts.sort(key=lambda r: r[0])
write_csv("posts.csv", ["投稿日時", "媒体", "投稿文", "いいね", "保存", "コメント数"], posts)
with open(os.path.join(OUT, "sample_posts.txt"), "w", encoding="utf-8") as fp:
    fp.write("過去の投稿の見本（練習用・架空）\n1つずつ空行で区切っています。自分の投稿に置き換えて使ってください。\n\n")
    fp.write("\n\n".join([
        "月末に慌てないコツは、入金と支払いの日を1枚の表に並べることです。むずかしい計算はいりません。まず並べる。それだけで、足りない日が見えてきます。 #資金繰り #中小企業",
        "銀行には、決算のときだけでなく毎月顔を出す。試算表を1枚持って、5分だけ。小さな習慣が、いざというときの話の早さになります。 #銀行対応",
        "「黒字なのにお金がない」。よく聞く言葉です。利益とお金は別もの。売掛金の回収日を、今日ひとつだけ確かめてみませんか。",
        "AIに頼むときは、目的・前提・形式の3つを書きます。「誰に向けて、何のために、どんな形で」。これだけで、返ってくる答えが変わります。 #AI活用",
        "今朝は現場まわりでした。数字の向こうに、働く人の顔があります。帰ってから見る試算表は、いつもより少しあたたかく見えました。",
        "前年同月と比べる。前月と比べるより、季節のクセが消えて、本当の変化が見えます。あなたの会社の9月は、去年とくらべてどうでしたか？ #数字の見方",
    ]) + "\n")

# ---------- 7. 外壁塗装の単価表 tanka-gaiheki.xlsx ----------
TANKA = [
    ("足場（架設・解体）", "㎡", 800), ("飛散防止ネット", "㎡", 150), ("高圧洗浄", "㎡", 200), ("養生", "㎡", 300),
    ("下地補修（ひび割れ補修）", "式", 30000), ("下塗り（シーラー）", "㎡", 700),
    ("中塗り・上塗り　シリコン塗料", "㎡", 2300), ("中塗り・上塗り　ラジカル塗料", "㎡", 2700), ("中塗り・上塗り　フッ素塗料", "㎡", 3800),
    ("シーリング打ち替え", "m", 900), ("付帯部塗装　雨どい", "m", 800), ("付帯部塗装　破風", "m", 900), ("付帯部塗装　軒天", "㎡", 1200),
    ("ベランダ防水（トップコート）", "㎡", 2500),
]
wb = Workbook()
ws = wb.active
ws.title = "単価表"
ws["A1"] = "外壁塗装 単価表（練習用・架空の単価です）"
ws["A1"].font = f(True, 14)
ws["A2"] = "実際の見積もりには、自社の単価表を使ってください。この数字は練習のための仮のものです。"
ws["A2"].font = f(False, 10, "5B6577")
header(ws, 4, ["項目", "単位", "単価（円・税抜）"])
for i, (item, unit, price) in enumerate(TANKA):
    r = 5 + i
    ws.cell(row=r, column=1, value=item)
    ws.cell(row=r, column=2, value=unit).alignment = Alignment(horizontal="center")
    ws.cell(row=r, column=3, value=price).number_format = "#,##0"
style_range(ws, 5, 4 + len(TANKA), 1, 3)
ws.cell(row=6 + len(TANKA), column=1, value="諸経費は、単価表に入れていません。「合計の5%」のように、お願い文で指定します。").font = f(False, 10, "5B6577")
base(ws, [34, 8, 18])
wb.save(os.path.join(OUT, "tanka-gaiheki.xlsx"))

# ---------- 集計の控え（検算用に表示） ----------
print("売上（店舗別）")
for ym in MONTHS:
    t = store_total(ym)
    print(" ", ym, {k: f"{v:,}" for k, v in t.items()}, "合計", f"{sum(t.values()):,}", "件数", len(SALES[ym]))
print("目標", targets)
d = {}
for r in SALES["2026-09"]:
    if r[2] == "デニムパンツ":
        d[r[1]] = d.get(r[1], 0) + r[7]
d2 = {}
for r in SALES["2025-09"]:
    if r[2] == "デニムパンツ":
        d2[r[1]] = d2.get(r[1], 0) + r[7]
print("デニム 当月", d, "前年", d2)
print("clients 明細合計(税抜)", sum(q * p for _, _, q, p, _ in LINES))
print("ABC 合計", sum(q * p for _, _, q, p in ABC))
